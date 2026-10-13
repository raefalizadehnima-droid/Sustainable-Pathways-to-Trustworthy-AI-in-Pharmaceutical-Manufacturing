#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Add a transparent template for four requested metrics to kham.xlsx.

Important:
The current workbook does not contain material masses, product yield,
energy consumption, or a defined comparison method. Therefore this script
does NOT fabricate PMI, PMI reduction, energy-normalized PMI, or partial r.
It preserves the existing data and adds an Additional_Metrics sheet showing
the required inputs and formulas/definitions.
"""

from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

INPUT_FILE = Path("kham.xlsx")
OUTPUT_FILE = Path("kham_with_additional_metrics_template.xlsx")

if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"Cannot find {INPUT_FILE}. Put this script in the same folder as kham.xlsx."
    )

data = pd.read_excel(INPUT_FILE)

required_existing = {
    "AI_Model_Accuracy (%)",
    "Batch_Failure_Rate (%)",
    "Production_Cost_Reduction (%)",
    "Regulatory_Compliance_Score",
    "System_Resilience_Score",
}
missing = required_existing.difference(data.columns)
if missing:
    raise ValueError(
        "The workbook does not have the expected columns: "
        + ", ".join(sorted(missing))
    )

metrics = pd.DataFrame([
    {
        "Requested metric": "PMI (Mean ± 95% CI)",
        "Result": "NOT CALCULABLE FROM CURRENT FILE",
        "Definition / calculation": (
            "PMI = total mass of materials used / mass of product obtained. "
            "Calculate the mean and t-based 95% CI across comparable runs."
        ),
        "Required additional data": (
            "Material input masses and product mass/yield for each run"
        ),
        "Reason / note": (
            "The existing columns describe AI and operational performance, "
            "not material inputs or product output."
        ),
    },
    {
        "Requested metric": "% Reduction",
        "Result": "NOT CALCULABLE FROM CURRENT FILE",
        "Definition / calculation": (
            "100 * (baseline PMI - comparison-method PMI) / baseline PMI"
        ),
        "Required additional data": (
            "PMI values for a clearly defined baseline and comparison method"
        ),
        "Reason / note": (
            "Production Cost Reduction (%) is not equivalent to PMI reduction."
        ),
    },
    {
        "Requested metric": "Energy-Normalized PMI",
        "Result": "NOT CALCULABLE FROM CURRENT FILE",
        "Definition / calculation": (
            "Specify the study's operational definition first. One possible "
            "definition is PMI divided by energy consumption, with units stated."
        ),
        "Required additional data": (
            "PMI and energy consumption for each comparable run, with units"
        ),
        "Reason / note": "The workbook contains no energy-consumption column.",
    },
    {
        "Requested metric": "Partial r (controlling yield)",
        "Result": "NOT CALCULABLE FROM CURRENT FILE",
        "Definition / calculation": (
            "Partial correlation between two specified variables after "
            "statistically controlling for yield."
        ),
        "Required additional data": (
            "Yield plus the two variables whose partial correlation is intended"
        ),
        "Reason / note": (
            "No yield column is present, and the two correlation variables "
            "must be identified."
        ),
    },
])

with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    data.to_excel(writer, sheet_name="Original_Data", index=False)
    metrics.to_excel(writer, sheet_name="Additional_Metrics", index=False)

wb = load_workbook(OUTPUT_FILE)
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col_cells in ws.columns:
        width = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col_cells[: min(ws.max_row, 150) + 1]
        )
        ws.column_dimensions[
            get_column_letter(col_cells[0].column)
        ].width = min(max(width + 2, 14), 60)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUTPUT_FILE)
print(f"Saved: {OUTPUT_FILE.resolve()}")
